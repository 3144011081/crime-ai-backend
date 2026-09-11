"""
Internship Progress Report Generator (July 15, 2026 – September 11, 2026)
Candidate: Iqra Rani (Employee ID: NEU000047)
Department: Software Development
Manager: Tayyaba Hussain
Email: iqra.rani@neuronixtech.net | Contact: 03700752043
Organization: Neuronix Technologies
Track: Artificial Intelligence & Computer Vision Engineering
Formatting: Times New Roman, 14pt Bold Headings, 12pt Subheadings, 11pt Body,
            Pure Black Font, Professional Monochrome Tables, Running Headers/Footers
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

doc = docx.Document()

# Page Margins (1.0 inch all around)
for section in doc.sections:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.different_first_page_header_footer = True

    # Running Header (Pages 2+)
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("Neuronix Technologies  |  Internship Progress Report — Iqra Rani (NEU000047)")
    hrun.font.name = "Times New Roman"
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(0, 0, 0)
    
    # Running Footer (Pages 2+)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    frun1 = fp.add_run("AI Aerial Surveillance & Crime Detection System		Page ")
    frun1.font.name = "Times New Roman"
    frun1.font.size = Pt(8.5)
    frun1.font.color.rgb = RGBColor(0, 0, 0)
    
    frun2 = fp.add_run()
    frun2.font.name = "Times New Roman"
    frun2.font.size = Pt(8.5)
    frun2.font.color.rgb = RGBColor(0, 0, 0)
    
    # Dynamic Page Number Field
    fld1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instr = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fld2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fld3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    frun2._r.append(fld1)
    frun2._r.append(instr)
    frun2._r.append(fld2)
    frun2._r.append(fld3)

FONT_NAME = "Times New Roman"
H1_SIZE = Pt(14)
H2_SIZE = Pt(12)
BODY_SIZE = Pt(11)

def _set_font(run, bold=False, italic=False, underline=False, size=None, color=None):
    run.bold = bold
    run.italic = italic
    run.underline = underline
    run.font.name = FONT_NAME
    run.font.size = size or BODY_SIZE
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    else:
        run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), FONT_NAME)
    rf.set(qn("w:hAnsi"), FONT_NAME)
    rf.set(qn("w:cs"), FONT_NAME)

def _set_spacing(para, before="60", after="40", line="280"):
    pPr = para._element.get_or_add_pPr()
    old = pPr.find(qn("w:spacing"))
    if old is not None:
        pPr.remove(old)
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), before)
    sp.set(qn("w:after"), after)
    sp.set(qn("w:line"), line)
    sp.set(qn("w:lineRule"), "auto")
    pPr.append(sp)

def _shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def _cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def _cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def _tbl_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def add_heading_1(text, page_break_before=False):
    p = doc.add_paragraph()
    if page_break_before:
        p.paragraph_format.page_break_before = True
    r = p.add_run(text)
    _set_font(r, bold=True, size=H1_SIZE, color=(0, 0, 0))
    _set_spacing(p, before="140", after="50", line="300")
    p.paragraph_format.keep_with_next = True
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_font(r, bold=True, size=H2_SIZE, color=(0, 0, 0))
    _set_spacing(p, before="100", after="30", line="280")
    p.paragraph_format.keep_with_next = True
    return p

def add_heading_3(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_font(r, bold=True, italic=True, size=Pt(11), color=(0, 0, 0))
    _set_spacing(p, before="70", after="20", line="250")
    p.paragraph_format.keep_with_next = True
    return p

def add_body(text, bold_prefix=None, indent=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.left_indent = Inches(0.3)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        _set_font(r1, bold=True, size=BODY_SIZE, color=(0, 0, 0))
        r2 = p.add_run(text)
        _set_font(r2, size=BODY_SIZE, color=(0, 0, 0))
    else:
        r = p.add_run(text)
        _set_font(r, size=BODY_SIZE, color=(0, 0, 0))
    _set_spacing(p, before="15", after="30", line="250")
    return p

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        _set_font(r1, bold=True, size=BODY_SIZE, color=(0, 0, 0))
        r2 = p.add_run(text)
        _set_font(r2, size=BODY_SIZE, color=(0, 0, 0))
    else:
        r = p.add_run(text)
        _set_font(r, size=BODY_SIZE, color=(0, 0, 0))
    _set_spacing(p, before="10", after="15", line="230")
    return p

def add_table(headers, rows, col_widths=None, is_toc=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_row = table.rows[0]
    _cant_split(hdr_row)
    _tbl_header(hdr_row)
    
    hdr_pad = 40 if is_toc else 50
    row_pad = 25 if is_toc else 40
    hdr_font_sz = Pt(9.5) if is_toc else Pt(10)
    row_font_sz = Pt(9) if is_toc else Pt(9.5)
    line_sp = "190" if is_toc else "220"
    
    # Headers
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = ""
        _shade_cell(cell, "D9D9D9")
        _cell_margins(cell, top=hdr_pad, bottom=hdr_pad, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        _set_font(r, bold=True, size=hdr_font_sz, color=(0, 0, 0))
        _set_spacing(p, before="8", after="8", line=line_sp)
        
    # Rows
    for r_idx, row_data in enumerate(rows):
        row = table.add_row()
        _cant_split(row)
        bg_color = "F5F5F5" if r_idx % 2 == 0 else "FFFFFF"
        for i, cell_text in enumerate(row_data):
            cell = row.cells[i]
            cell.text = ""
            _shade_cell(cell, bg_color)
            _cell_margins(cell, top=row_pad, bottom=row_pad, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(cell_text))
            _set_font(r, size=row_font_sz, color=(0, 0, 0))
            _set_spacing(p, before="8", after="8", line=line_sp)
            
    if col_widths:
        for i, w in enumerate(col_widths):
            for r in table.rows:
                r.cells[i].width = Inches(w)
                
    return table

def add_log_table(rows, col_widths=[1.1, 1.4, 4.0]):
    headers = ["Day", "Date", "Tasks Performed & Technical Deliverables"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_row = table.rows[0]
    _cant_split(hdr_row)
    _tbl_header(hdr_row)
    
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = ""
        _shade_cell(cell, "D9D9D9")
        _cell_margins(cell, top=50, bottom=50, left=70, right=70)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        _set_font(r, bold=True, size=Pt(10), color=(0, 0, 0))
        _set_spacing(p, before="6", after="6", line="210")
        
    for r_idx, (day_str, date_str, tasks) in enumerate(rows):
        row = table.add_row()
        _cant_split(row)
        bg_color = "F5F5F5" if r_idx % 2 == 0 else "FFFFFF"
        is_special = day_str in ["Sick Leave", "Casual Leave", "Holiday"]
        
        # Col 0: Day
        cell_day = row.cells[0]
        cell_day.text = ""
        _shade_cell(cell_day, bg_color)
        _cell_margins(cell_day, top=40, bottom=40, left=50, right=50)
        p_day = cell_day.paragraphs[0]
        p_day.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_day = p_day.add_run(day_str)
        _set_font(r_day, bold=True, italic=is_special, size=Pt(9.5), color=(0, 0, 0))
        _set_spacing(p_day, before="6", after="6", line="210")
        
        # Col 1: Date
        cell_date = row.cells[1]
        cell_date.text = ""
        _shade_cell(cell_date, bg_color)
        _cell_margins(cell_date, top=40, bottom=40, left=50, right=50)
        p_date = cell_date.paragraphs[0]
        p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_date = p_date.add_run(date_str)
        _set_font(r_date, bold=False, size=Pt(9.5), color=(0, 0, 0))
        _set_spacing(p_date, before="6", after="6", line="210")
        
        # Col 2: Tasks
        cell_tasks = row.cells[2]
        cell_tasks.text = ""
        _shade_cell(cell_tasks, bg_color)
        _cell_margins(cell_tasks, top=40, bottom=40, left=80, right=80)
        
        if isinstance(tasks, str):
            p = cell_tasks.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(tasks)
            _set_font(r, italic=True, size=Pt(9.5), color=(0, 0, 0))
            _set_spacing(p, before="6", after="6", line="210")
        else:
            for t_idx, (b_pfx, t_text) in enumerate(tasks):
                p = cell_tasks.paragraphs[0] if t_idx == 0 else cell_tasks.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                if b_pfx:
                    rb = p.add_run(f"• {b_pfx} ")
                    _set_font(rb, bold=True, size=Pt(9.5), color=(0, 0, 0))
                else:
                    rb = p.add_run("• ")
                    _set_font(rb, bold=True, size=Pt(9.5), color=(0, 0, 0))
                rt = p.add_run(t_text)
                _set_font(rt, size=Pt(9.5), color=(0, 0, 0))
                _set_spacing(p, before="2", after="3", line="200")
                
    if col_widths:
        for i, w in enumerate(col_widths):
            for r in table.rows:
                r.cells[i].width = Inches(w)
                
    return table


# ══════════════════════════════════════════════════════════════════════════════
# PAGE 1: FORMAL EXECUTIVE COVER PAGE
# ══════════════════════════════════════════════════════════════════════════════

# Organization Header
p_org = doc.add_paragraph()
p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_org = p_org.add_run("NEURONIX TECHNOLOGIES")
_set_font(r_org, bold=True, size=Pt(14), color=(0, 0, 0))
_set_spacing(p_org, before="40", after="20", line="260")

p_dept = doc.add_paragraph()
p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_dept = p_dept.add_run("Department of Software Development & Artificial Intelligence")
_set_font(r_dept, bold=False, italic=True, size=Pt(11), color=(0, 0, 0))
_set_spacing(p_dept, before="0", after="100", line="240")

# Report Title
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_title = p_title.add_run("INTERNSHIP PROGRESS REPORT")
_set_font(r_title, bold=True, size=Pt(18), color=(0, 0, 0))
_set_spacing(p_title, before="40", after="30", line="340")

p_proj = doc.add_paragraph()
p_proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_proj = p_proj.add_run("AI Aerial Surveillance & Crime Detection System from Drone (UAV) Video Feeds")
_set_font(r_proj, bold=True, italic=True, size=Pt(11.5), color=(0, 0, 0))
_set_spacing(p_proj, before="0", after="80", line="260")

# Candidate Profile Table
add_table(
    ["Particular", "Detail"],
    [
        ["Candidate / Intern Name", "Iqra Rani"],
        ["Employee ID", "NEU000047"],
        ["Designation / Track", "Intern — Artificial Intelligence & Computer Vision Engineering"],
        ["Department", "Software Development"],
        ["Manager / Supervisor", "Tayyaba Hussain (Project Lead)"],
        ["Host Organization", "Neuronix Technologies"],
        ["Official Email", "iqra.rani@neuronixtech.net | Contact: 03700752043"],
        ["Internship Tenure", "July 15, 2026 – September 11, 2026 (9 Weeks / 38 Working Days)"],
        ["Reporting Status", "Completed / Full Technical Milestones, Modernization & Cloud Production Deployment"],
        ["Key Handover Deliverables", "Live Vercel React 18 SPA, Cloudflare GPU Tunnel, Dual GitHub Repositories, Project Technical Report ('Project_Report_Crime_Detection_System-1'), & GUI Video ('project gui.mov')"]
    ],
    col_widths=[2.4, 4.1]
)

# Formal Endorsement Sign-off on Cover Page
p_auth = doc.add_paragraph()
p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_auth = p_auth.add_run("Verified & Submitted for Supervisory Endorsement — Neuronix Technologies")
_set_font(r_auth, bold=True, italic=True, size=Pt(9.5), color=(0, 0, 0))
_set_spacing(p_auth, before="80", after="0", line="220")


# ══════════════════════════════════════════════════════════════════════════════
# PAGE 2: TABLE OF CONTENTS
# ══════════════════════════════════════════════════════════════════════════════
add_heading_1("Table of Contents", page_break_before=True)

add_table(
    ["Section #", "Section Title", "Key Content Covered"],
    [
        ["1", "Intern Introduction & Project Context", "Candidate profile overview, credentials, technical competencies, and assigned project scope"],
        ["2", "Phase-wise Technical Work Summary", "Phases 1 to 8: onboarding, ML theory, cloud AI, dataset engineering, bias resolution, 2-stage AI, React 18 modernization, and Vercel/Cloudflare deployment"],
        ["3", "Comprehensive Day-by-Day Work Log (Mon–Fri)", "Detailed daily technical log table covering all working days (Days 1 to 38) across Weeks 1 to 9, including leaves and holidays (till September 11)"],
        ["3.1", "  • Week 1: Onboarding, Environment & Tooling", "Daily log table: workstation setup, macOS MPS acceleration, Git/GitHub, Docker, and approved sick leave (Jul 17)"],
        ["3.2", "  • Week 2: ML Foundations & Video Classification", "Daily log table: ML evaluation metrics, CNN backbones, R(2+1)D convolution factorization, YOLOv8 tracking"],
        ["3.3", "  • Week 3: Cloud AI Platforms (Colab & Kaggle) & Remote Training", "Daily log table: Cloud GPU runtimes (NVIDIA T4), Kaggle API & dataset ingestion, Colab drive mounting, remote PyTorch workflows"],
        ["3.4", "  • Week 4: Aerial Dataset & Baseline 3D CNN", "Daily log table: UCF-Crime curation, aerial perspective warping, data loaders, baseline CrimeR2Plus1D model"],
        ["3.5", "  • Week 5: Training & Class Bias Diagnosis", "Daily log table: Training Run 1, Robbery bias discovery, class-weighted loss, casual leave (Aug 13), Independence Day"],
        ["3.6", "  • Week 6: 2-Stage Hybrid AI Architecture", "Daily log table: Pivot to 2-stage framework, drone person detector, specialized weapon detector, ByteTrack, latency benchmark"],
        ["3.7", "  • Week 7: Preprocessing, Zoom & Super-Resolution", "Daily log table: crime_aerial_augmented_best1.pth, 5-mode preprocessor, SceneQualityMetrics, FSRCNN 3x zoom"],
        ["3.8", "  • Week 8: Strict Gating & Flask Web Portal", "Daily log table: False shooting alarm fix, strict 6-rule gating, tensor ghost fix, Flask backend, MJPEG live stream & ngrok remote testing"],
        ["3.9", "  • Week 9: React 18 SPA, Vercel & Cloudflare Tunnel", "Daily log table: React 18 + Vite migration, dynamic incident chart, monoserver split, Vercel cloud deploy, Cloudflare GPU tunnel & final handover"],
        ["4", "Three AI Models Specification & Handover Deliverables", "Architecture specs table (Crime 3D CNN, Drone Person YOLO, Weapon YOLO), gating matrix, live deployment links, and formal deliverables"],
        ["5", "Conclusion & Supervisory Verification Sign-Off", "Professional competencies gained, formal performance evaluation sheet, and supervisory sign-off"]
    ],
    col_widths=[1.0, 2.6, 2.9],
    is_toc=True
)


# ══════════════════════════════════════════════════════════════════════════════
# PAGE 3: SECTION 1: INTERN INTRODUCTION & PROJECT CONTEXT
# ══════════════════════════════════════════════════════════════════════════════
add_heading_1("1. Intern Introduction & Project Context", page_break_before=True)

add_body(
    "My name is Iqra Rani (Employee ID: NEU000047), working as an Intern in the Software Development "
    "Department at Neuronix Technologies under the managerial supervision and technical mentorship of Tayyaba Hussain. "
    "My primary focus during this engineering tenure is centered on research, deep learning model development, "
    "and full-stack software system implementation in computer vision, spatio-temporal video action classification, multi-object "
    "detection and tracking, edge latency profiling, and production web deployment."
)

add_heading_2("1.1 Technical Specialization & Core Competencies")
add_body(
    "My technical domain covers applied deep learning, neural network optimization, and computer vision systems. "
    "Throughout the project lifecycle, I specialized in PyTorch development (leveraging Apple Silicon Metal Performance "
    "Shaders - MPS and CUDA), Ultralytics YOLOv8 object detection architectures, ByteTrack persistent multi-object tracking, 3D Convolutional "
    "Neural Networks (R(2+1)D-18), image enhancement algorithms (CLAHE, dark channel dehazing, AGCWD gamma correction, FSRCNN super-resolution), "
    "edge tensor optimization (FP16 half-precision), modular monoserver architecture (Flask REST blueprints), modern Single Page Applications (React 18, Vite, TypeScript, Tailwind CSS), "
    "cloud continuous delivery (Vercel), and hardware-accelerated remote inference pipelines (Cloudflare Tunnels)."
)

add_heading_2("1.2 Assigned Project Scope & Core Responsibilities")
add_body(
    "Under the managerial direction of Tayyaba Hussain at Neuronix Technologies, I was assigned as the lead developer for the "
    "'AI Aerial Surveillance & Crime Detection System from Drone (UAV) Video Feeds'. My core responsibilities included:"
)
add_bullet("Designing and implementing an aerial video domain adaptation and augmentation pipeline (`aerial_augmentation_solution.py`).", "Data Engineering: ")
add_bullet("Training, diagnosing, and fine-tuning 3D ResNet/R(2+1)D action recognition models on UCF-Crime and aerial-adapted video datasets.", "Model Training: ")
add_bullet("Diagnosing and successfully resolving complex machine learning class imbalance biases (Robbery and Shooting over-prediction).", "Bias Resolution: ")
add_bullet("Integrating spatial object detection (Drone Person Detector + Specialized Weapon Detector) to construct a robust 2-Stage Hybrid AI architecture.", "Multi-Model Fusion: ")
add_bullet("Developing a strict 6-rule Multi-Modal Crime Decision Engine (`fuse_multimodal_crime_decision`) to eliminate false shooting alerts on normal scenes.", "Decision Engine: ")
add_bullet("Building an Adaptive Video Acquisition system (`adaptive_acquisition.py`) with multi-scale ROI zoom cropping and FSRCNN 3x super-resolution.", "Adaptive Acquisition: ")
add_bullet("Architecting a modular Python monoserver backend (`backend/monoserver.py`) and decoupling into a modern React 18 + Vite SPA frontend with synchronized dual-view video playback.", "Modern Full-Stack: ")
add_bullet("Deploying production web application to Vercel (https://aerial-crime-detector.vercel.app) and bridging Apple Silicon MPS GPU acceleration via Cloudflare Quick Tunnel for global access.", "Cloud Deployment: ")
add_bullet("Benchmarking edge latency, memory profiles, and multi-altitude drone viewpoints for robust aerial surveillance readiness.", "Edge Optimization: ")


# ══════════════════════════════════════════════════════════════════════════════
# PAGE 4: SECTION 2: PHASE-WISE TECHNICAL WORK SUMMARY
# ══════════════════════════════════════════════════════════════════════════════
add_heading_1("2. Phase-wise Technical Work Summary", page_break_before=True)

add_heading_2("Phase 1: Onboarding, DevOps & Tooling")
add_body("Configured local macOS workstation with PyTorch Metal Performance Shaders (MPS) hardware acceleration, established GitFlow version control workflows, and authored base Dockerfile for containerized inference.")

add_heading_2("Phase 2: Machine Learning & Video Classification Theory")
add_body("Studied supervised ML loss dynamics, evaluated spatio-temporal action recognition architectures comparing 2D CNNs against 3D R(2+1)D factorization, and analyzed YOLOv8 with ByteTrack persistent tracking.")

add_heading_2("Phase 3: Cloud AI Platforms (Colab & Kaggle) & Cloud Training")
add_body("Configured cloud GPU runtimes (NVIDIA T4 on Colab & Kaggle), automated dataset ingestion via Kaggle API, mounted persistent Google Drive storage, and benchmarked remote PyTorch data loaders and training loops.")

add_heading_2("Phase 4: Aerial Dataset Engineering & Baseline 3D CNN")
add_body("Curated UCF-Crime benchmark, developed aerial perspective warping and CLAHE augmentation pipeline (`aerial_augmentation_solution.py`), and constructed baseline CrimeR2Plus1D PyTorch action classifier.")

add_heading_2("Phase 5: Resolving Class Imbalance & Predictive Bias")
add_body("Diagnosed Robbery and Shooting over-prediction biases; addressed feature ambiguity through balanced batch sampling, loss re-weighting, and recognized necessity of spatial evidence gating.")

add_heading_2("Phase 6: 2-Stage Multi-Modal Pipeline & Adaptive Zoom")
add_body("Architected 2-Stage Hybrid AI framework coupling spatial YOLOv8 detectors (drone person, weapons, vehicles) with 3D CNN temporal classification and FSRCNN 3x super-resolution adaptive zoom.")

add_heading_2("Phase 7: Strict Evidence Gating & Video HUD Polish")
add_body("Engineered strict 6-rule Multi-Modal Crime Decision Engine achieving 0% false shooting alarms, enforced 95% confidence cap, streamlined top surveillance HUD overlays, and verified automated evidence extraction.")

add_heading_2("Phase 8: Enterprise Modernization, React 18 SPA Migration & Cloud Production Deployment")
add_body("Transitioned system from monolithic rendering to a decoupled microservices architecture. Constructed a component-driven React 18 + Vite + TypeScript frontend with synchronized dual-view video playback and pure dynamic audit analytics. Decoupled into two GitHub repositories (`-crime-ai-frontend` and `crime-ai-backend`), deployed frontend live to Vercel (https://aerial-crime-detector.vercel.app), and bridged Apple Silicon MPS GPU acceleration to the web via Cloudflare Tunnel (https://reserved-inf-pod-trends.trycloudflare.com).")


# ══════════════════════════════════════════════════════════════════════════════
# SECTION 3: COMPLETE DAY-BY-DAY (MONDAY TO FRIDAY) DETAILED WORK LOG
# ══════════════════════════════════════════════════════════════════════════════
add_heading_1("3. Comprehensive Day-by-Day Detailed Work Log (Monday to Friday)", page_break_before=False)
add_body(
    "The following section provides the exact day-by-day record of all engineering activities, research breakthroughs, "
    "code implementations, bug resolutions, and testing milestones achieved across every working day from Monday through Friday across all 9 weeks."
)

# ─── WEEK 1 (15 Jul – 17 Jul 2026) ───
add_heading_2("3.1 Week 1: Onboarding, Environment Configuration & Tooling (15 Jul – 17 Jul 2026)")
add_log_table([
    (
        "Day 1",
        "Wed, Jul 15, 2026",
        [
            ("Orientation:", "Completed formal internship onboarding at Neuronix Technologies with Manager Tayyaba Hussain; received project goals and technical roadmap."),
            ("Environment Setup:", "Configured local macOS workstation with Python 3.13 virtual environment (`venv`) and package dependencies."),
            ("Hardware Check:", "Verified PyTorch Metal Performance Shaders (MPS) device acceleration for hardware-accelerated local deep learning."),
            ("Repository:", "Initialized project repository and established project directory layout and documentation guidelines.")
        ]
    ),
    (
        "Day 2",
        "Thu, Jul 16, 2026",
        [
            ("Git & GitHub:", "Practiced version control workflows: branching models (GitFlow), atomic commit guidelines, and pull requests."),
            ("Repository Config:", "Configured `.gitignore` rules for large binary checkpoint files (.pth, .pt) and dataset storage."),
            ("Docker Tooling:", "Studied containerization concepts and authored base `Dockerfile` configuring OpenCV, PyTorch, TorchVision, and headless video codecs."),
            ("Validation:", "Conducted initial code review of team baseline scripts and verified local container build and inference execution.")
        ]
    ),
    (
        "Sick Leave",
        "Fri, Jul 17, 2026",
        "Approved medical sick leave availed due to indisposition."
    )
])

# ─── WEEK 2 (20 Jul – 24 Jul 2026) ───
add_heading_2("3.2 Week 2: Machine Learning Foundations & Video Classification Theory (20 Jul – 24 Jul 2026)")
add_log_table([
    (
        "Day 3",
        "Mon, Jul 20, 2026",
        [
            ("ML Foundations:", "Reviewed classical machine learning theory: supervised vs. unsupervised paradigms and classification loss functions."),
            ("Evaluation Metrics:", "Studied evaluation metrics critical for imbalanced crime detection: Precision, Recall, F1-Score, and ROC-AUC."),
            ("Optimization:", "Investigated gradient descent optimizers (SGD with momentum vs. Adam vs. AdamW with cosine learning rate scheduling).")
        ]
    ),
    (
        "Holiday",
        "Tue, Jul 21, 2026",
        "Scheduled Public / Company Holiday observed."
    ),
    (
        "Day 4",
        "Wed, Jul 22, 2026",
        [
            ("Deep Learning:", "Analyzed Deep Convolutional Neural Networks (CNNs) for visual feature extraction: ResNet residual skip connections."),
            ("Transfer Learning:", "Studied transfer learning techniques: freezing backbone weights vs. end-to-end full fine-tuning on downstream tasks."),
            ("Backbone Study:", "Explored TorchVision model zoo implementations and pre-trained Kinetic-400 video backbone weights.")
        ]
    ),
    (
        "Day 5",
        "Thu, Jul 23, 2026",
        [
            ("Video Architectures:", "Investigated spatio-temporal video architectures: comparing 2D CNN + LSTM against 3D CNNs (C3D, I3D, R3D)."),
            ("R(2+1)D Theory:", "Detailed study of the R(2+1)D (Residual (2+1)D) convolution factorization (1x3x3 spatial + 3x1x1 temporal)."),
            ("Complexity Analysis:", "Evaluated computational complexity and temporal window requirements (clip frame length, sliding window stride).")
        ]
    ),
    (
        "Day 6",
        "Fri, Jul 24, 2026",
        [
            ("Object Detection:", "Explored real-time object detection: YOLOv8 architecture (CSPDarknet backbone, PANet neck, anchor-free decoupled head)."),
            ("Tracking Theory:", "Studied Non-Maximum Suppression (NMS) and persistent multi-object tracking using ByteTrack Kalman filtering."),
            ("Pipeline Design:", "Prepared architectural feasibility comparison for combining spatial object tracking with 3D temporal classification.")
        ]
    )
])

# ─── WEEK 3 (27 Jul – 31 Jul 2026) ───
add_heading_2("3.3 Week 3: Cloud AI Platforms (Google Colab & Kaggle) & Cloud Training (27 Jul – 31 Jul 2026)")
add_log_table([
    (
        "Day 7",
        "Mon, Jul 27, 2026",
        [
            ("Colab Onboarding:", "Onboarded onto Google Colab cloud development environment; evaluated runtime types (Standard CPU vs High-RAM with GPU)."),
            ("GPU Configuration:", "Configured GPU acceleration (NVIDIA T4 / V100) and verified PyTorch CUDA backend availability and driver compatibility."),
            ("Cloud Storage:", "Configured persistent storage integration by mounting Google Drive (`google.colab.drive`) for storing weights and checkpoints.")
        ]
    ),
    (
        "Day 8",
        "Tue, Jul 28, 2026",
        [
            ("Kaggle Exploration:", "Explored Kaggle cloud notebooks and GPU compute infrastructure (dual NVIDIA T4 instances with 30-hour weekly quota)."),
            ("Kaggle API Setup:", "Generated and configured Kaggle API credentials (`kaggle.json`) on local and cloud instances for secure command-line automation."),
            ("CLI Automation:", "Tested CLI automation scripts (`kaggle datasets download`) to streamline rapid dataset synchronization to cloud environments.")
        ]
    ),
    (
        "Day 9",
        "Wed, Jul 29, 2026",
        [
            ("Dataset Loading:", "Practiced loading, streaming, and extracting large surveillance video datasets directly onto Colab and Kaggle NVMe drives."),
            ("I/O Profiling:", "Addressed cloud I/O bottlenecks: compared zipped archive extraction speeds against uncompressed video streaming across runtimes."),
            ("Storage Optimization:", "Engineered optimized dataset unarchiving routines using `shutil` and `tarfile` to ensure immediate read-readiness on ephemeral disks.")
        ]
    ),
    (
        "Day 10",
        "Thu, Jul 30, 2026",
        [
            ("Script Porting:", "Ported PyTorch neural network training pipelines into modular Jupyter notebooks (`.ipynb`) tailored for Colab and Kaggle execution."),
            ("CUDA Tuning:", "Configured GPU memory allocation flags (`torch.cuda.empty_cache()`, `cudnn.benchmark = True`) to maximize cloud batch throughput."),
            ("Loader Optimization:", "Tested multi-worker `DataLoader` configurations to eliminate CPU preprocessing starvation during heavy video decoding.")
        ]
    ),
    (
        "Day 11",
        "Fri, Jul 31, 2026",
        [
            ("Prototype Training:", "Executed prototype multi-epoch deep learning training runs on sample video batches using Kaggle and Colab GPU runtimes."),
            ("Cloud Checkpointing:", "Implemented automated checkpoint synchronization to Google Drive and Kaggle Working directories at every epoch completion."),
            ("Weekly Review:", "Compiled a comprehensive Cloud Training Guide & SOP for the engineering team under Manager Tayyaba Hussain.")
        ]
    )
])

# ─── WEEK 4 (03 Aug – 07 Aug 2026) ───
add_heading_2("3.4 Week 4: Aerial Dataset Engineering & Baseline 3D CNN Pipeline (03 Aug – 07 Aug 2026)")
add_log_table([
    (
        "Day 12",
        "Mon, Aug 03, 2026",
        [
            ("Project Inception:", "Official project kick-off: 'AI Aerial Surveillance & Crime Detection System from Drone Video Feeds'."),
            ("Taxonomy:", "Formulated core project taxonomy: Normal, Violence, Robbery, Shooting, FireExplosion, Accident, Vandalism."),
            ("Gap Analysis:", "Investigated available open-source crime video datasets and documented the absence of dedicated aerial UAV benchmarks.")
        ]
    ),
    (
        "Day 13",
        "Tue, Aug 04, 2026",
        [
            ("Dataset Curation:", "Downloaded and organized the UCF-Crime ground-level surveillance benchmark dataset."),
            ("Data Profiling:", "Analyzed video metadata, resolution variance, codec incompatibilities, and extreme class distribution skew."),
            ("Data Splitting:", "Created video validation splits ensuring scene independence across training and testing partitions.")
        ]
    ),
    (
        "Day 14",
        "Wed, Aug 05, 2026",
        [
            ("Preprocessing Engine:", "Engineered video preprocessing pipeline: OpenCV frame decoders, Kinetics channel normalization, 112x112 resizing."),
            ("Data Loaders:", "Constructed PyTorch `Dataset` and `DataLoader` classes supporting dynamic sliding window clip extraction."),
            ("Performance Test:", "Tested frame loading throughput and memory footprints under multi-threaded worker configurations.")
        ]
    ),
    (
        "Day 15",
        "Thu, Aug 06, 2026",
        [
            ("Data Augmentation:", "Authored `aerial_augmentation_solution.py` to bridge the domain gap between ground-level CCTV and aerial UAV footage."),
            ("Perspective Warping:", "Implemented geometric perspective transforms: top-down nadir warping, high-altitude scale reduction, and rotation."),
            ("Visual Filters:", "Integrated environmental filters: atmospheric haze simulation, random motion blur, and CLAHE contrast equalization.")
        ]
    ),
    (
        "Day 16",
        "Fri, Aug 07, 2026",
        [
            ("Model Architecture:", "Implemented the baseline `CrimeR2Plus1D` PyTorch model architecture in `models/crime_model.py`."),
            ("Head Design:", "Added classification head with Dropout (p=0.5) and 7 output class logits."),
            ("Unit Testing:", "Verified model instantiation, tensor forward-pass shapes [1, 3, 8, 112, 112] -> [1, 7], and initial loss computation.")
        ]
    )
])

# ─── WEEK 5 (10 Aug – 14 Aug 2026) ───
add_heading_2("3.5 Week 5: Model Training, Class Bias Diagnosis & Sampling Experiments (10 Aug – 14 Aug 2026)")
add_log_table([
    (
        "Day 17",
        "Mon, Aug 10, 2026",
        [
            ("Training Run 1:", "Initiated baseline training of CrimeR2Plus1D on UCF-Crime dataset using Adam optimizer and Cross-Entropy loss."),
            ("Monitoring:", "Logged epoch-by-epoch loss convergence, validation accuracy, and macro F1 scores."),
            ("Checkpointing:", "Saved initial model checkpoint `crime_r2plus1d_ucf_finetuned.pth` (355 MB).")
        ]
    ),
    (
        "Day 18",
        "Tue, Aug 11, 2026",
        [
            ("Bias Discovery:", "Conducted inference evaluation on validation drone footage; uncovered critical **Class Bias Issue 1 (Robbery Bias)**."),
            ("Diagnosis:", "The model heavily over-predicted Robbery across normal walking scenes due to disproportionate Robbery sample counts."),
            ("Quantitative Audit:", "Quantified false positive distribution using per-class confusion matrix analysis in `bias_evaluation.py`.")
        ]
    ),
    (
        "Day 19",
        "Wed, Aug 12, 2026",
        [
            ("Class Weighting:", "Formulated Fine-tuning Solution 1: Weighted Cross-Entropy Loss to penalize over-represented crime classes."),
            ("Loss Modification:", "Calculated inverse class frequency weights and integrated them into the PyTorch training criterion."),
            ("Balanced Sampler:", "Implemented balanced random sampling per batch in `DataLoader` to address minority class bias without gradient destabilization.")
        ]
    ),
    (
        "Casual Leave",
        "Thu, Aug 13, 2026",
        "Approved casual leave availed for personal commitments."
    ),
    (
        "Holiday",
        "Fri, Aug 14, 2026",
        "Independence Day National Holiday observed."
    )
])

# ─── WEEK 6 (17 Aug – 21 Aug 2026) ───
add_heading_2("3.6 Week 6: 2-Stage Hybrid Architecture & YOLOv8 Integration (17 Aug – 21 Aug 2026)")
add_log_table([
    (
        "Day 20",
        "Mon, Aug 17, 2026",
        [
            ("Architecture Pivot:", "Synthesized training findings: spatial physical evidence is mandatory to prevent temporal action false positives."),
            ("System Design:", "Architected the **2-Stage Hybrid AI Framework**: Stage 1 Spatial Entity Tracking + Stage 2 Temporal Classification."),
            ("Interface Design:", "Drafted interface specifications for passing detector telemetry into temporal decision logic.")
        ]
    ),
    (
        "Day 21",
        "Tue, Aug 18, 2026",
        [
            ("Person Detector:", "Integrated custom aerial fine-tuned YOLOv8 person detector `drone_person_detector_best.pt` (18.3 MB)."),
            ("Threshold Tuning:", "Configured detection threshold (conf >= 0.20) optimized for high-altitude nadir and oblique human silhouettes."),
            ("Tracking Integration:", "Coupled detector with ByteTrack persistent tracker to compute per-person trajectory and speed metrics (px/frame).")
        ]
    ),
    (
        "Day 22",
        "Wed, Aug 19, 2026",
        [
            ("Weapon Detector:", "Integrated specialized weapon detector `weapon_detector_best.pt` (18.3 MB) supporting Firearms, Blades, and Explosives."),
            ("Crop Search:", "Implemented contextual high-resolution person bounding-box crop searching to catch small aerial weapon signatures."),
            ("Vehicle Tracker:", "Added general COCO YOLOv8s tracker for vehicle entity verification (cars, trucks, motorcycles, buses).")
        ]
    ),
    (
        "Day 23",
        "Thu, Aug 20, 2026",
        [
            ("Detector Class:", "Built the unified `ObjectActivityDetector` class in `video_inference1.py` integrating all three YOLO detectors."),
            ("Interaction Engine:", "Implemented spatial proximity and bounding-box intersection logic for `people_are_interacting()` algorithm."),
            ("Integration Check:", "Verified telemetry contracts between Stage 1 tracking objects and downstream decision logic.")
        ]
    ),
    (
        "Day 24",
        "Fri, Aug 21, 2026",
        [
            ("Benchmarking:", "Benchmarked Stage 1 multi-detector latency on Apple MPS GPU, verifying sub-30ms per-frame execution."),
            ("Stress Testing:", "Stress-tested detector pipeline on crowded aerial footage with variable drone camera angles."),
            ("Weekly Audit:", "Compiled Week 6 multi-detector performance audit report for Manager Tayyaba Hussain.")
        ]
    )
])

# ─── WEEK 7 (24 Aug – 28 Aug 2026) ───
add_heading_2("3.7 Week 7: Image Preprocessing, Adaptive Zoom & Super-Resolution (24 Aug – 28 Aug 2026)")
add_log_table([
    (
        "Day 25",
        "Mon, Aug 24, 2026",
        [
            ("Model Upgrade:", "Trained and integrated upgraded primary crime model `crime_aerial_augmented_best1.pth` (358 MB, 224 state tensors)."),
            ("Calibration:", "Implemented post-training Softmax Temperature Scaling (T=0.7) to produce properly calibrated probability distributions."),
            ("Model Loader:", "Updated unified loader `load_crime_model()` in `models/crime_model.py` with automatic checkpoint inspection.")
        ]
    ),
    (
        "Day 26",
        "Tue, Aug 25, 2026",
        [
            ("Preprocessor:", "Authored `image_preprocessing.py` (836 lines) providing scene-adaptive visual enhancement pipelines."),
            ("Filter Modes:", "Implemented 5 preprocessing modes: `auto`, `aerial_drone` (CLAHE + dark channel dehazing), `low_light_night` (AGCWD), `yolo_enhanced`, and `none`."),
            ("Performance Check:", "Validated preprocessor throughput reaching 50.42 FPS (19.83 ms/frame) on Apple MPS.")
        ]
    ),
    (
        "Day 27",
        "Wed, Aug 26, 2026",
        [
            ("Quality Metrics:", "Developed `SceneQualityMetrics` class analyzing brightness, RMS contrast, Laplacian sharpness, and colorfulness."),
            ("Auto-Routing:", "Implemented automated scene diagnosis routing low-contrast or hazy drone footage to appropriate filter pipelines."),
            ("Testing:", "Created unit tests in `test_preprocessing.py` and visualization comparisons in `visualize_preprocessing.py`.")
        ]
    ),
    (
        "Day 28",
        "Thu, Aug 27, 2026",
        [
            ("Adaptive Acquisition:", "Authored `adaptive_acquisition.py` (329 lines) resolving high-altitude small-target action ambiguity."),
            ("Zoom Strategy:", "Implemented dynamic multi-scale ROI zoom cropping testing scale factors [1.5x, 2.0x, 2.5x, 3.0x]."),
            ("Super-Resolution:", "Integrated FSRCNN (Fast Super-Resolution CNN) 3x upscale model (`models/FSRCNN_x3.pb`) for zoomed ROI clarity.")
        ]
    ),
    (
        "Day 29",
        "Fri, Aug 28, 2026",
        [
            ("Fusion Logic:", "Drafted initial multi-modal reasoning engine connecting detector telemetry with 3D CNN temporal predictions."),
            ("Adaptive Clarity:", "Implemented `predict_clip_with_adaptive_clarity()` triggering CLAHE re-evaluation on low confidence (<0.48)."),
            ("Pipeline Test:", "Conducted full end-to-end pipeline dry-runs across diverse test video clips.")
        ]
    )
])

# ─── WEEK 8 (31 Aug – 04 Sep 2026) ───
add_heading_2("3.8 Week 8: Strict Evidence Gating, False Alarm Elimination & Flask App (31 Aug – 04 Sep 2026)")
add_log_table([
    (
        "Day 30",
        "Mon, Aug 31, 2026",
        [
            ("Bug Investigation:", "Conducted systematic false alarm audit; identified that walking scenes were falsely triggering Shooting classifications."),
            ("Root Cause:", "Identified root causes: unconstrained temporal sum bias and lack of hard physical weapon gating."),
            ("System Redesign:", "Redesigned multi-modal fusion architecture to enforce physical evidence prerequisites before any crime is asserted.")
        ]
    ),
    (
        "Day 31",
        "Tue, Sep 01, 2026",
        [
            ("Strict Gating:", "Rewrote `fuse_multimodal_crime_decision()` in `video_inference1.py` with **Strict 6-Rule Priority Gating** (Rules 1 to 6)."),
            ("Verification:", "Validated all 7 gating verification test scenarios with 100% pass rate.")
        ]
    ),
    (
        "Holiday",
        "Wed, Sep 02, 2026",
        "Scheduled Public Sector Holiday observed."
    ),
    (
        "Day 32",
        "Thu, Sep 03, 2026",
        [
            ("Tensor Fix:", "Fixed ghost probability contamination in sliding window loop by zeroing tensor and storing only active labels."),
            ("Video Aggregation:", "Calibrated video-level crime confirmation rule requiring `crime_ratio >= 0.15` or top crime confidence >= 0.40."),
            ("Evidence Extraction:", "Implemented automated evidence clip extraction (`extract_evidence_clips()`) generating standalone MP4 clips.")
        ]
    ),
    (
        "Day 33",
        "Fri, Sep 04, 2026",
        [
            ("Flask Backend:", "Developed production Flask Web Application (`app.py`, 262 lines) supporting video uploads and sample selection."),
            ("Frontend UI:", "Constructed responsive Jinja2 HTML/CSS templates: Dashboard (`index.html`), Live Feed (`live.html`), Reports (`reports.html`)."),
            ("Live Streaming:", "Implemented real-time MJPEG live streaming route `/live_video/<filename>` for in-browser surveillance monitoring."),
            ("Ngrok Remote Inference:", "Configured ngrok reverse tunneling (`pyngrok`) to expose the local Flask inference server to a secure public HTTPS endpoint; verified remote video inference streaming.")
        ]
    )
])

# ─── WEEK 9 (07 Sep – 11 Sep 2026) ───
add_heading_2("3.9 Week 9: Confidence Calibration, UI Polishing & Final Handover (07 Sep – 11 Sep 2026)")
add_log_table([
    (
        "Day 34",
        "Mon, Sep 07, 2026",
        [
            ("Confidence Cap:", "Enforced confidence score capping rule: all crime predictions capped at a maximum of 95% (`min(conf, 0.95)`)."),
            ("Logic Update:", "Applied confidence capping across `predict_clip()`, `fuse_multimodal_crime_decision()`, and video-level summaries."),
            ("UI Upgrade:", "Upgraded Flask dashboard UI (`templates/index.html`) with bold red/green Crime Alert Boxes and animated progress bars.")
        ]
    ),
    (
        "Day 35",
        "Tue, Sep 08, 2026",
        [
            ("HUD Refinement:", "Refined HUD video overlays in `video_inference1.py`: removed misaligned weapon bounding boxes in favor of clean scene boxes."),
            ("Visual Polish:", "Removed noisy subtext line from top video HUD banner and streamlined banner height to 54px for sleek surveillance visuals."),
            ("Visual Audit:", "Verified annotated video outputs and evidence clip generation with updated HUD graphics.")
        ]
    ),
    (
        "Day 36",
        "Wed, Sep 09, 2026",
        [
            ("System Audit:", "Conducted comprehensive system verification on Apple MPS hardware; validated inference speed (1.47x real-time)."),
            ("Batch Evaluation:", "Executed batch evaluation suite `evaluate_model_videos.py` across UCF-Crime and aerial test splits."),
            ("Metrics Audit:", "Logged performance metrics: 89.3% accuracy on normal walking scenes and 0% false Shooting rate on unarmed scenes.")
        ]
    ),
    (
        "Day 37",
        "Thu, Sep 10, 2026",
        [
            ("Documentation:", "Authored automated DOCX project report generator `generate_project_report.py` using `python-docx`."),
            ("System Specs:", "Formatted complete system architecture diagrams, model layer breakdowns, and Flask route tables."),
            ("Audit Schema:", "Verified automated JSON forensic report schemas and evidence clip chain-of-custody metadata storage in `outputs/`.")
        ]
    ),
    (
        "Day 38",
        "Fri, Sep 11, 2026",
        [
            ("React 18 SPA Migration:", "Modernized user interface from legacy Jinja2 templates into a responsive, component-driven React 18 + Vite + TypeScript + Tailwind CSS application featuring dark mode, Lucide icons, and live status badges."),
            ("Dual-View Video Player:", "Engineered synchronized dual-feed video player allowing side-by-side comparison of raw aerial surveillance feeds against AI-annotated crime detection streams."),
            ("Dynamic Audit Analytics:", "Completely refactored the Incident Distribution Chart and Chronological Audit Log from static values into 100% dynamic calculations computed directly from `window_audits` in the inference engine."),
            ("Monoserver Architecture:", "Refactored backend into a modular Monoserver architecture (`backend/monoserver.py`, `backend/config.py`, REST API blueprints in `backend/api/`) supporting videos, inference, live MJPEG feeds, reports, and settings."),
            ("Dual-Repository Split:", "Decoupled codebase into two production GitHub repositories: Frontend (`3144011081/-crime-ai-frontend`) and Backend (`3144011081/crime-ai-backend`)."),
            ("Vercel Cloud Deployment:", "Deployed modern React frontend to Vercel global CDN with custom domain alias: `https://aerial-crime-detector.vercel.app`, configuring SPA rewrites and caching policies in `vercel.json`."),
            ("Cloudflare GPU Tunnel:", "Deployed Cloudflare Quick Tunnel (`cloudflared`) bridging local Apple Silicon Metal Performance Shaders (MPS) hardware acceleration to the public internet (`https://reserved-inf-pod-trends.trycloudflare.com`), enabling zero-cost, zero-card, full-speed remote inference."),
            ("Cloud Containerization:", "Enhanced `Dockerfile` with lightweight PyTorch CPU wheels and graceful checkpoint fallback to support deployment across multi-cloud environments (Render, Koyeb, Spaces)."),
            ("Final Demo & Handover:", "Conducted comprehensive final project presentation to Manager Tayyaba Hussain, presenting live Vercel web application, Cloudflare GPU tunnel, 2-stage multi-modal AI gating, and formally delivering the Project Technical Report and GUI demonstration video."),
            ("Tenure Conclusion:", "Successfully concluded the 9-week engineering internship tenure at Neuronix Technologies (September 11, 2026).")
        ]
    )
])

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 4: THREE AI MODELS SPECIFICATION & HANDOVER DELIVERABLES
# ══════════════════════════════════════════════════════════════════════════════
add_heading_1("4. Three AI Models Specification & Handover Deliverables", page_break_before=True)

add_body(
    "The final production system successfully integrates three specialized, coordinated AI models resident in GPU/MPS memory:"
)

add_table(
    ["Model Name", "Model File", "Size", "Architecture", "Primary Role in Pipeline"],
    [
        ["Crime Classifier", "crime_aerial_augmented_best1.pth", "358 MB", "CrimeR2Plus1D (3D CNN)", "Classifies 7 spatio-temporal crime actions over 8-frame clips."],
        ["Drone Person Detector", "drone_person_detector_best.pt", "18.3 MB", "YOLOv8 Custom Aerial", "Detects humans from UAV angles (conf >= 0.20) + ByteTrack speed tracking."],
        ["Weapon Detector", "weapon_detector_best.pt", "18.3 MB", "YOLOv8 Specialized", "Detects Firearms, Blades & Explosives (conf >= 0.15) for physical gating."]
    ],
    col_widths=[1.5, 1.8, 0.8, 1.5, 1.6]
)

add_heading_2("4.1 Multi-Modal Decision Gating Verification Matrix")
add_table(
    ["Scenario / Physical State", "Detector Inputs", "Expected Verdict", "System Output", "Status"],
    [
        ["Separated Walkers", "2 People, No Overlap, No Weapons", "Normal", "Normal (92.0%)", "PASSED ✓"],
        ["Single Walker", "1 Person, Low Speed, No Weapons", "Normal", "Normal (92.0%)", "PASSED ✓"],
        ["Physical Confrontation", "2 People, Overlapping Bounding Boxes", "Violence", "Violence (70.0%)", "PASSED ✓"],
        ["Snatch-and-Run", "Interaction + High Speed Fleeing", "Robbery", "Robbery (70.0%)", "PASSED ✓"],
        ["Shooting Pred. (No Gun)", "Model says Shooting, No Firearm", "Normal", "Normal (92.0%)", "PASSED ✓"],
        ["Armed Shooting Threat", "Model says Shooting + Firearm Found", "Shooting", "Shooting (75.0%)", "PASSED ✓"],
        ["Armed Blade Assault", "Blade / Knife Detected", "Violence", "Violence (75.0%)", "PASSED ✓"]
    ],
    col_widths=[1.4, 1.9, 1.0, 1.2, 0.9]
)

add_heading_2("4.2 Formal Project Handover Deliverables")
add_body(
    "On Friday, September 11, 2026, the formal engineering handover was successfully conducted with Manager Tayyaba Hussain. "
    "The primary project handover deliverables submitted include:"
)
add_table(
    ["Deliverable Item", "File Name / Reference", "Description & Scope Handed Over"],
    [
        ["Live Production Web App", "https://aerial-crime-detector.vercel.app", "Production-grade React 18 + Vite + TypeScript SPA deployed live on Vercel with global CDN, dual video players, real-time audit logs, and dynamic analytics."],
        ["Hardware-Accelerated Tunnel", "https://reserved-inf-pod-trends.trycloudflare.com", "Cloudflare Quick Tunnel bridging Apple Silicon MPS (Metal Performance Shaders) GPU acceleration to the web for zero-cost, high-speed remote video inference."],
        ["Frontend Git Repository", "https://github.com/3144011081/-crime-ai-frontend", "Complete modern frontend codebase with TypeScript definitions, Tailwind styling, and Vercel build configuration."],
        ["Backend Git Repository", "https://github.com/3144011081/crime-ai-backend", "Complete Python Monoserver repository with REST API blueprints, multi-stage Dockerfile, and cloud deployment configs."],
        ["Project Technical Report", "Project_Report_Crime_Detection_System-1", "Comprehensive technical report handed over to Manager Tayyaba Hussain covering system design, 3D CNN & YOLOv8 architectures, aerial domain augmentation, strict 6-rule gating, and empirical benchmark evaluations."],
        ["Project GUI & Functioning Video", "project gui.mov", "High-definition video demonstration exhibiting all GUI features, web dashboard controls, real-time MJPEG live detection streaming, evidence clip extraction, and full end-to-end system functioning."],
        ["Crime Classification Model", "crime_aerial_augmented_best1.pth (358 MB)", "Production-ready 3D CNN (R(2+1)D-18) checkpoint fine-tuned on aerial-augmented surveillance data with calibrated Softmax temperature scaling."],
        ["Drone Person Detection Model", "drone_person_detector_best.pt (18.3 MB)", "Specialized YOLOv8 aerial detector checkpoint trained for high-altitude nadir and oblique human silhouette localization with ByteTrack trajectory tracking."],
        ["Specialized Weapon Detection Model", "weapon_detector_best.pt (18.3 MB)", "Dedicated YOLOv8 detector trained for localized firearm, knife/blade, and explosive identification for physical evidence gating."],
        ["Inference Suite & Forensic Schemas", "Monoserver (`backend/`), Engine (`video_inference1.py`)", "Complete inference pipeline, multi-format forensic reports (JSON, DOCX), and evidence clip chain-of-custody logging."]
    ],
    col_widths=[1.7, 2.1, 3.2]
)

# ══════════════════════════════════════════════════════════════════════════════
# SECTION 5: CONCLUSION & SUPERVISORY VERIFICATION SIGN-OFF
# ══════════════════════════════════════════════════════════════════════════════
add_heading_1("5. Conclusion & Supervisory Verification Sign-Off", page_break_before=True)
add_body(
    "The engineering tenure at Neuronix Technologies under the mentorship and supervision of Tayyaba Hussain provided comprehensive "
    "hands-on experience spanning computer vision research, deep learning model architecture design, addressing real-world class imbalance, "
    "multi-modal reasoning engine development, edge optimization, modern React 18 frontend engineering, and cloud deployment. The key professional outcomes achieved include:"
)
add_bullet("Mastery of spatio-temporal deep learning architectures (R(2+1)D, 3D convolutions, sliding window inference).", "Deep Learning: ")
add_bullet("Expertise in YOLOv8 fine-tuning, multi-object tracking (ByteTrack), and multi-detector orchestration.", "Computer Vision: ")
add_bullet("Designing domain adaptation strategies to successfully bridge ground-level CCTV data to aerial drone perspectives.", "Domain Adaptation: ")
add_bullet("Eliminating machine learning bias and hallucinations through strict multi-modal physical evidence gating.", "Bias Resolution: ")
add_bullet("Leveraging cloud GPU platforms (Google Colab, Kaggle) for rapid dataset loading, distributed experimentation, and high-performance PyTorch training.", "Cloud AI & Training: ")
add_bullet("Architecting decoupled full-stack systems: React 18 + Vite SPA, Python Monoserver REST APIs, and Vercel cloud continuous deployment.", "Modern Web Architecture: ")
add_bullet("Bridging Apple Silicon Metal Performance Shaders (MPS) hardware acceleration to global web clients via Cloudflare Quick Tunnels.", "Edge & Hybrid Cloud: ")

add_heading_2("5.1 Formal Supervisory Verification & Sign-Off")
add_body(
    "This progress report and all associated project deliverables (including live cloud deployments, technical documentation, "
    "source repositories, and model checkpoints) have been reviewed and approved by the assigned project supervisor:"
)

# Formal Sign-off Sheet Table
add_table(
    ["Verification Attribute", "Candidate / Intern Submission", "Supervisory Review & Endorsement"],
    [
        ["Full Name", "Iqra Rani", "Tayyaba Hussain"],
        ["Employee ID / Designation", "NEU000047 | Intern (Software Dev)", "Project Lead / Manager (Software Dev)"],
        ["Organization", "Neuronix Technologies", "Neuronix Technologies"],
        ["Technical Milestone Verdict", "All 9 Weeks Completed / Milestones Achieved", "Verified & Approved / Criteria Fully Satisfied"],
        ["Deliverables Received", "Live Vercel App, Cloudflare Tunnel, Dual Repos, Project Report, & GUI Video", "Confirmed Received & Validated in Working Order"],
        ["Signature & Date", "_______________________  [Sep 11, 2026]", "_______________________  [Sep 11, 2026]"]
    ],
    col_widths=[2.1, 2.2, 2.2]
)

# Save to destination
output_path = "/Users/ranaasad/ML_Project copy/Internship_Progress_Report.docx"
doc.save(output_path)
print(f"[OK] Successfully generated: {output_path}")
print(f"     File size: {os.path.getsize(output_path) / 1024:.1f} KB")
